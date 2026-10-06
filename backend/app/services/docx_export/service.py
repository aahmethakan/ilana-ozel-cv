from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.services.public_cv_projection import PublicCvProjection


class DocxExportError(ValueError): pass


def export_reviewed_draft(projection: PublicCvProjection) -> bytes:
    """Render only the reviewed public draft; no evidence or generation logic lives here."""
    reviewed = projection.reviewed_draft
    review = reviewed.review
    if not (review.skills or review.work_entries or review.education):
        raise DocxExportError("Reviewed draft has no exportable content.")
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10)
    title = document.add_paragraph(projection.name or "CV")
    title.style = document.styles["Title"]; title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if projection.headline: document.add_paragraph(projection.headline)
    contact = " | ".join(value for value in (projection.contact.email, projection.contact.phone, projection.contact.website) if value)
    if contact: document.add_paragraph(contact)
    def heading(text: str): document.add_heading(text, level=1)
    if projection.summary:
        heading("PROFESSIONAL SUMMARY")
        document.add_paragraph(projection.summary)
    if review.skills:
        heading("SKILLS")
        document.add_paragraph(", ".join(dict.fromkeys(item.text for item in review.skills)))
    if review.work_entries:
        heading("WORK EXPERIENCE")
        for entry in review.work_entries:
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{entry.title} | {entry.company}").bold = True
            if entry.dates: document.add_paragraph(entry.dates)
            for claim in entry.claims: document.add_paragraph(claim.text, style="List Bullet")
    if review.education:
        heading("EDUCATION")
        for entry in review.education:
            text = entry.qualification or entry.institution
            if entry.qualification: text = f"{text} | {entry.institution}"
            if entry.dates: text = f"{text} | {entry.dates}"
            document.add_paragraph(text)
    stream = BytesIO(); document.save(stream)
    data = stream.getvalue()
    validate_docx_export(data, projection)
    return data


def validate_docx_export(data: bytes, projection: PublicCvProjection) -> None:
    if not data:
        raise DocxExportError("DOCX export is empty.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist(): raise DocxExportError("DOCX structure is invalid.")
            text = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    except DocxExportError: raise
    except Exception as error: raise DocxExportError("DOCX structure is invalid.") from error
    forbidden = ("evidence_id", "association_id", "candidate_id", "session_id", "draft:")
    if any(value in text for value in forbidden): raise DocxExportError("DOCX contains internal metadata.")
    for item in projection.reviewed_draft.review.skills:
        if item.text not in text: raise DocxExportError("DOCX did not preserve reviewed skills.")
