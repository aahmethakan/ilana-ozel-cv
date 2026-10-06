from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches, Pt

from app.services.public_cover_letter_projection import PublicCoverLetterProjection


class CoverLetterDocxExportError(ValueError):
    pass


_FORBIDDEN = ("session_id", "draft_id", "item_id", "evidence_id", "association_evidence_id", "atomic_evidence_id", "candidate_id", "work_record_id", "work_candidate_id", "verification_status", "source_type", "job_requirement_id", "match_id")


def export_cover_letter_docx(projection: PublicCoverLetterProjection) -> bytes:
    if not (projection.opening.strip() and projection.closing.strip()):
        raise CoverLetterDocxExportError("Cover letter has no exportable content.")
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)
    if projection.candidate_name:
        document.add_paragraph(projection.candidate_name).runs[0].bold = True
    contact = " | ".join(value for value in (projection.contact.email, projection.contact.phone, projection.contact.website) if value)
    if contact:
        document.add_paragraph(contact)
    if projection.target_role:
        target = f"Application for {projection.target_role}"
        if projection.target_company:
            target += f" at {projection.target_company}"
        document.add_paragraph(target)
    document.add_paragraph("Dear Hiring Manager,")
    document.add_paragraph(projection.opening)
    for paragraph in projection.body_sections:
        document.add_paragraph(paragraph)
    document.add_paragraph(projection.closing)
    document.add_paragraph("Kind regards,")
    if projection.candidate_name:
        document.add_paragraph(projection.candidate_name)
    stream = BytesIO(); document.save(stream)
    data = stream.getvalue()
    validate_cover_letter_docx(data, projection)
    return data


def validate_cover_letter_docx(data: bytes, projection: PublicCoverLetterProjection) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            text = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception as error:
        raise CoverLetterDocxExportError("DOCX structure is invalid.") from error
    if not text or any(value in text.lower() for value in _FORBIDDEN):
        raise CoverLetterDocxExportError("DOCX is empty or contains internal metadata.")
    for value in (projection.candidate_name, projection.target_role, projection.target_company, projection.opening, *projection.body_sections, projection.closing):
        if value and value not in text:
            raise CoverLetterDocxExportError("DOCX did not preserve reviewed public content.")
